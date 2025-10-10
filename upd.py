import streamlit as st
import qrcode
import zipfile
import os
from io import BytesIO
import socket
import urllib.parse
from openpyxl import Workbook, load_workbook


# Descobre o IP local
def get_local_ip():
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
    finally:
        s.close()
    return ip


def gerar_xlsx(local, numeros_serie):
    xlsx_path = os.path.abspath(f"qrcode_files/{local}.xlsx")
    pasta = os.path.dirname(xlsx_path)
    os.makedirs(pasta, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Números de Série"

    for i, numero in enumerate(numeros_serie, start=1):
        ws.cell(row=i, column=1, value=numero)

    wb.save(xlsx_path)
    return xlsx_path


def gerar_qrCodes_zip(local, base_url):
    zip_buffer = BytesIO()
    encoded_local = urllib.parse.quote(local)
    internal_url = f"{base_url}/?file={encoded_local}"

    with zipfile.ZipFile(zip_buffer, "w") as zipf:
        qr = qrcode.QRCode(
            version=1,
            error_correction=qrcode.constants.ERROR_CORRECT_H,
            box_size=10,
            border=4,
        )
        qr.add_data(internal_url)
        qr.make(fit=True)

        img = qr.make_image(fill_color="black", back_color="white")
        img_bytes = BytesIO()
        img.save(img_bytes, format="PNG")
        img_bytes.seek(0)

        zipf.writestr(f"{os.path.basename(local)}.png", img_bytes.read())

    zip_buffer.seek(0)
    return zip_buffer, internal_url


def app():
    with open("assets/header.html", "r", encoding="utf-8") as f:
        st.markdown(f.read(), unsafe_allow_html=True)

    try:
        query_params = st.query_params
    except AttributeError:
        query_params = st.experimental_get_query_params()

    if "file" in query_params:
        local_name = (
            query_params["file"]
            if isinstance(query_params["file"], str)
            else query_params["file"][0]
        )
        local_name = urllib.parse.unquote(local_name)
        file_path = os.path.abspath(f"qrcode_files/{local_name}.xlsx")

        if os.path.exists(file_path):
            wb = load_workbook(file_path)
            ws = wb.active
            conteudo = "\n".join(
                [
                    str(cell)
                    for row in ws.iter_rows(values_only=True)
                    for cell in row
                    if cell
                ]
            )

            st.subheader(f"Números de série - {local_name}")
            st.text(conteudo)

            with open(file_path, "rb") as f:
                st.download_button(
                    "Baixar arquivo Excel",
                    data=f,
                    file_name=f"{local_name}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )

        else:
            st.error("Arquivo não encontrado.")
        return  # interrompe o fluxo principal

    # Interface principal
    local = st.text_input("Local do Estoque:", placeholder="Ex.: PRIME/Disponível")
    entrada = st.text_area(
        "Insira os números de série:", placeholder="Digite um número de série por linha"
    )

    if st.button("Gerar QR Code"):
        numeros = [n.strip() for n in entrada.splitlines() if n.strip()]

        if numeros:
            gerar_xlsx(local, numeros)

            ip = get_local_ip()
            base_url = f"http://{ip}:8501"

            zip_buffer, url = gerar_qrCodes_zip(local, base_url)

            st.success("QR Code gerado com sucesso!")
            st.write(f"Acesse os números de série pelo link: [Abrir arquivo]({url})")

            st.download_button(
                label="Baixar QR Code",
                data=zip_buffer,
                file_name="qrcodes.zip",
                mime="application/zip",
            )
        else:
            st.warning("Insira pelo menos um número de série.")

    # Estilo e rodapé
    with open("assets/style.css") as f:
        st.markdown(f"<style>{f.read()}</style>", unsafe_allow_html=True)

    with open("assets/footer.html", encoding="utf-8") as f:
        st.markdown(f.read(), unsafe_allow_html=True)
